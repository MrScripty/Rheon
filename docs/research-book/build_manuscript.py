#!/usr/bin/env python3
"""Build a native-equation Word book from chapter Markdown."""
from pathlib import Path
from copy import deepcopy
import json,re,subprocess,zipfile
from lxml import etree
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
R=Path(__file__).resolve().parent;O=R/'output';O.mkdir(exist_ok=True)
B=R/'build';B.mkdir(exist_ok=True)
chapters=sorted((R/'chapters').glob('*.md'));appendices=sorted((R/'appendices').glob('*.md'));files=chapters+appendices
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
formulas=[]
for f in files:
 for m in re.finditer(r'\\\[(.*?)\\\]|\\\((.*?)\\\)',f.read_text(encoding="utf-8"),re.S):
  x=(m.group(1) or m.group(2)).strip()
  if x not in formulas:formulas.append(x)
(B/'equations.md').write_text('\n\n'.join(f'MATHINDEX{i:04d}\n\n$$\n{x}\n$$' for i,x in enumerate(formulas)),encoding='utf-8')
subprocess.run(['pandoc','-f','markdown+tex_math_dollars','-t','docx',str(B/'equations.md'),'-o',str(B/'equations.docx')],check=True)
with zipfile.ZipFile(B/'equations.docx') as z:body=etree.fromstring(z.read('word/document.xml'))
equations={};current=None
for p in body.findall('.//w:body/w:p',NS):
 t=''.join(p.xpath('.//w:t/text()',namespaces=NS))
 if t.startswith('MATHINDEX'):current=int(t[9:])
 else:
  nodes=p.findall('.//m:oMath',NS)
  if nodes and current is not None:equations[formulas[current]]=deepcopy(nodes[0]);current=None
assert set(formulas)==set(equations)
doc=Document();doc.core_properties.title='Rheon Discrete Fluid Simulation'
doc.core_properties.author='Puma';doc.core_properties.last_modified_by='Puma'
doc.core_properties.subject='Mathematics algorithms implementation and checked contracts'
s=doc.styles['Normal'];s.font.name='Liberation Serif';s.font.size=Pt(10.5)
s.paragraph_format.line_spacing=1.12;s.paragraph_format.space_after=Pt(6);s.paragraph_format.widow_control=True
for name,size in [('Title',28),('Subtitle',13),('Heading 1',21),('Heading 2',13)]:
 s=doc.styles[name];s.font.name='DejaVu Sans';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
 s.font.bold=name.startswith('Heading');s.paragraph_format.keep_with_next=True
 s.paragraph_format.space_before=Pt(12);s.paragraph_format.space_after=Pt(7);s.paragraph_format.line_spacing=1.05
doc.styles['Heading 1'].paragraph_format.page_break_before=True
doc.styles['Caption'].font.name='Liberation Serif';doc.styles['Caption'].font.size=Pt(9)
doc.styles['Caption'].font.bold=False;doc.styles['Caption'].font.color.rgb=RGBColor(45,55,60)
for style in doc.styles:
 for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
def geometry(s):
 s.page_width=Inches(7);s.page_height=Inches(10.5)
 s.top_margin=s.bottom_margin=Inches(.67);s.left_margin=Inches(.72);s.right_margin=Inches(.65)
 s.header_distance=s.footer_distance=Inches(.3)
def background(path):
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(1)
 run=p.add_run();run.font.size=Pt(1)
 inline=run.add_picture(str(path),width=Inches(7),height=Inches(10.5))._inline
 a=OxmlElement('wp:anchor')
 for k,v in [('distT','0'),('distB','0'),('distL','0'),('distR','0'),('simplePos','0'),('relativeHeight','0'),
             ('behindDoc','1'),('locked','0'),('layoutInCell','1'),('allowOverlap','1')]:a.set(k,v)
 sp=OxmlElement('wp:simplePos');sp.set('x','0');sp.set('y','0');a.append(sp)
 for d in ['H','V']:
  pos=OxmlElement('wp:position'+d);pos.set('relativeFrom','page')
  off=OxmlElement('wp:posOffset');off.text='0';pos.append(off);a.append(pos)
 a.append(deepcopy(inline.find(qn('wp:extent'))));a.append(OxmlElement('wp:wrapNone'))
 for n in ['wp:docPr','wp:cNvGraphicFramePr','a:graphic']:
  child=inline.find(qn(n))
  if child is not None:a.append(deepcopy(child))
 inline.getparent().replace(inline,a)
def footer(s,number=True):
 s.footer.is_linked_to_previous=False;p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 if number:
  el=OxmlElement('w:fldSimple');el.set(qn('w:instr'),'PAGE');p._p.append(el)
def bookmark(p,name,index):
 a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(index));a.set(qn('w:name'),name)
 b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),str(index));p._p.insert(0,a);p._p.append(b)
def link(p,text,url=None,anchor=None):
 el=OxmlElement('w:hyperlink')
 if url:el.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
 if anchor:el.set(qn('w:anchor'),anchor)
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'225D7A')
 pr.append(c);r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);el.append(r);p._p.append(el)
def mathnode(tex):
 node=deepcopy(equations[tex.strip()])
 # Normalize schema order: LibreOffice otherwise misreads closing delimiters.
 for pr in node.findall('.//m:dPr',NS):
  children=list(pr)
  for child in children:pr.remove(child)
  for name in ['begChr','sepChr','endChr','grow','shp','ctrlPr']:
   for child in children:
    if etree.QName(child).localname==name:pr.append(child)
 for r in node.findall('.//m:r',NS):
  pr=r.find(qn('w:rPr'))
  if pr is None:pr=OxmlElement('w:rPr');r.insert(0,pr)
  fonts=OxmlElement('w:rFonts')
  for k in ['ascii','hAnsi','cs']:fonts.set(qn('w:'+k),'DejaVu Math TeX Gyre')
  pr.append(fonts);sz=OxmlElement('w:sz');sz.set(qn('w:val'),'21');pr.append(sz)
 return node
token=re.compile(r'\\\(.*?\\\)|\[[^\]]+\]\([^)]+\)',re.S)
def inline(p,text):
 last=0
 for m in token.finditer(text):
  p.add_run(text[last:m.start()]);v=m.group()
  if v.startswith(r'\('):p._p.append(mathnode(v[2:-2]))
  else:
   label,url=re.match(r'\[([^\]]+)\]\(([^)]+)\)',v).groups();link(p,label,url=url)
  last=m.end()
 p.add_run(text[last:])
figures={
 '03-staggered-grids.md':[('Divergence from flux balance','staggered-grid','A staggered two-dimensional slice; z velocity occupies the corresponding third-axis faces.')],
 '05-linear-solvers.md':[('Residual and warm starts','pressure-residual','Executed residuals on the gauge-fixed 8 cubed graph. Iterations are not equal-cost units.'),('V cycle specification','multigrid-mechanism','Executed one-dimensional Dirichlet two-level mechanism, not obstacle-aware three-dimensional multigrid.')],
 '06-transport.md':[('Row sums and column sums','interpolation-mass','An algebraic row-stochastic counterexample changes total scalar while preserving bounds.'),('Conservative flux updates','box-diffusion','Periodic constant translation conserves total in this special case but diffuses shape.'),('Trace integrators','backtrace-refinement','Executed rotating-field trajectory errors for two backtrace integrators.')],
 '07-corrected-advection.md':[('Executed refinement','transport-refinement','Executed sine-wave refinement at Courant number one half, including the stated limiter.')],
 '09-liquid-surfaces.md':[('Surface tension time scales','curvature-refinement','Executed sphere curvature stencil error, not a moving-droplet validation.')],
 '13-memory-and-parallelism.md':[('Ownership and access','memory-scaling','Calculated named-field memory. Overhead and several production resources are excluded.')],
 '15-floating-point.md':[('Invalid values and finite range','rounding-gap','Executed arithmetic discrepancy. Binary32 uses rounded binary64 pressure, not a new solve.')],
 'E-worked-matrix-laboratory.md':[('Deriving energy reduction','cycle-circulation','A graph cycle retains zero-divergence circulation; it is an algebraic fixture.')]
}
count=0
def figure(name,caption):
 global count
 count+=1;p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
 shape=p.add_run().add_picture(str(R/'figures'/(name+'.png')),width=Inches(4.85 if name=='multigrid-mechanism' else 5.45))
 shape._inline.docPr.set('descr',caption)
 doc.add_paragraph(f'Figure {count}. {caption}',style='Caption')
geometry(doc.sections[0]);background(R/'artwork/front-cover.png');footer(doc.sections[0],False)
s=doc.add_section(WD_SECTION.NEW_PAGE);geometry(s);footer(s)
doc.add_paragraph('Rheon Discrete Fluid Simulation',style='Title')
doc.add_paragraph('Mathematics algorithms and checked discrete contracts',style='Subtitle')
doc.add_paragraph('Puma');doc.add_paragraph('Research edition October 2026')
doc.add_paragraph('A technical foundation for low-memory three-dimensional fluid simulation and geometric image guidance. This book derives discrete operators, transport, boundaries, surfaces and particle methods, then specifies a staged Rust implementation with explicit validation requirements.')
doc.add_paragraph('Includes 20 checked Lean theorems, reproducible numerical mechanism experiments and original explanatory figures. Exact discrete proofs do not establish production software, physical calibration or real-time performance.')
link(doc.add_paragraph(),'Research repository and qualification',url='https://github.com/MrScripty/Rheon/pull/1')
doc.add_paragraph('Contents',style='Heading 1')
for i,f in enumerate(files):
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3);link(p,f.read_text(encoding="utf-8").splitlines()[0][2:],anchor=f'section{i}')
for index,f in enumerate(files):
 chapter_start=len(doc.paragraphs)
 lines=f.read_text(encoding="utf-8").splitlines();i=0
 while i<len(lines):
  line=lines[i].strip()
  if not line:i+=1;continue
  if line.startswith('# '):
   p=doc.add_paragraph(line[2:],style='Heading 1');bookmark(p,f'section{index}',index+1);i+=1;continue
  if line.startswith('##'):
   title=line.lstrip('#').strip();doc.add_paragraph(title,style='Heading 2')
   for trigger,name,caption in figures.get(f.name,[]):
    if title==trigger:figure(name,caption)
   i+=1;continue
  if line==r'\[':
   chunk=[];i+=1
   while lines[i].strip()!=r'\]':chunk.append(lines[i]);i+=1
   p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_together=True
   p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(7)
   mp=OxmlElement('m:oMathPara');mp.append(mathnode('\n'.join(chunk)));p._p.append(mp);i+=1;continue
  if line.startswith('~~~'):
   chunk=[];i+=1
   while i<len(lines) and not lines[i].strip().startswith('~~~'):chunk.append(lines[i]);i+=1
   for j,code in enumerate(chunk):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
    p.paragraph_format.keep_with_next=j<len(chunk)-1 and len(chunk)<24;p.paragraph_format.left_indent=Inches(.1)
    r=p.add_run(code);r.font.name='Liberation Mono';r.font.size=Pt(8)
   doc.add_paragraph().paragraph_format.space_after=Pt(2);i+=1;continue
  if line.startswith('- '):
   p=doc.add_paragraph(style='List Bullet');inline(p,line[2:]);i+=1;continue
  chunk=[line];i+=1
  while i<len(lines) and lines[i].strip() and not lines[i].startswith(('#','- ','~~~',r'\[')):
   chunk.append(lines[i].strip());i+=1
  p=doc.add_paragraph();inline(p,' '.join(chunk))
 if f.name.startswith(('12-','14-','17-','18-')):
  for paragraph in doc.paragraphs[chapter_start:]:
   if paragraph.style.name in ('Normal','List Bullet'):
    paragraph.paragraph_format.space_after=Pt(3 if f.name.startswith('17-') else 2)
    if f.name.startswith(('12-','18-')):paragraph.paragraph_format.line_spacing=1.0
   elif paragraph.style.name=='Heading 2':
    paragraph.paragraph_format.space_before=Pt(9)
    paragraph.paragraph_format.space_after=Pt(5)
s=doc.add_section(WD_SECTION.NEW_PAGE);geometry(s);footer(s,False);background(R/'artwork/back-cover.png')
for text,size,bold,space in [
 ('A fluid solver built from explicit contracts',20,True,18),
 ('From staggered geometry and pressure projection to bounded transport, free surfaces, viscous stress and particle transfers, Rheon develops the mathematics behind a compact simulation framework.',12,False,15),
 ('Worked matrix examples, reproducible numerical experiments and checked Lean proofs connect exact identities to the practical limits of floating-point implementations.',12,False,12),
 ('A technical research book for low-memory three-dimensional fluids and reliable spatial guidance.',12,False,12)]:
 p=doc.add_paragraph();p.paragraph_format.space_before=Pt(space);p.paragraph_format.space_after=Pt(5)
 r=p.add_run(text);r.font.name='DejaVu Sans';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor(225,238,245)
doc.save(O/'rheon-discrete-fluid-simulation.docx')
combined='# Rheon Discrete Fluid Simulation\n\nAuthor Puma\n\n'+'\n\n'.join(f.read_text(encoding="utf-8") for f in files)
(O/'rheon-manuscript.md').write_text(combined,encoding="utf-8")
stats={'chapters':len(chapters),'appendices':len(appendices),'words':len(combined.split()),'unique_native_equations':len(equations),'figures':count}
(B/'assembly-stats.json').write_text(json.dumps(stats,indent=2)+'\n',encoding='utf-8');print(json.dumps(stats,indent=2))
